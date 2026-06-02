"""Resume / job text extraction from PDF, DOCX, and TXT files."""
from __future__ import annotations

import io
import os

import fitz  # PyMuPDF
import pdfplumber
from docx import Document

from ..config import settings


class FileValidationError(ValueError):
    pass


# Magic-byte signatures used as a lightweight content sniff. This is NOT a
# substitute for a real AV scan (see scan_for_malware) but rejects obvious
# mismatches between extension and content.
_SIGNATURES = {
    ".pdf": [b"%PDF"],
    ".docx": [b"PK\x03\x04"],  # docx is a zip container
}


def validate_file(filename: str, content: bytes) -> str:
    """Validate extension, size, and magic bytes. Returns the lowercase ext."""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in settings.ALLOWED_EXTENSIONS:
        raise FileValidationError(f"Unsupported file type: {ext}")
    if len(content) == 0:
        raise FileValidationError("Empty file")
    if len(content) > settings.MAX_UPLOAD_BYTES:
        raise FileValidationError("File exceeds maximum allowed size")
    sigs = _SIGNATURES.get(ext)
    if sigs and not any(content.startswith(s) for s in sigs):
        raise FileValidationError(f"File content does not match {ext} signature")
    return ext


def scan_for_malware(content: bytes) -> None:
    """Hook for malware scanning.

    In production, stream `content` to an AV engine (e.g. ClamAV via clamd, or
    a cloud scanning API) and raise FileValidationError on a positive hit.
    Left as an integration point so the rest of the pipeline is wired for it.
    """
    # Example ClamAV wiring (requires a running clamd):
    #   import clamd; cd = clamd.ClamdNetworkSocket(); r = cd.instream(io.BytesIO(content))
    #   if r["stream"][0] == "FOUND": raise FileValidationError("Malware detected")
    return None


def _extract_pdf(content: bytes) -> str:
    text_parts: list[str] = []
    # pdfplumber is best for layout-aware text; PyMuPDF is the fallback.
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                text_parts.append(page.extract_text() or "")
        text = "\n".join(text_parts).strip()
        if text:
            return text
    except Exception:
        pass
    try:
        with fitz.open(stream=content, filetype="pdf") as doc:
            return "\n".join(page.get_text() for page in doc).strip()
    except Exception:
        return ""


def _extract_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


def _extract_txt(content: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return content.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore").strip()


def extract_text(filename: str, content: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(content)
    if ext == ".docx":
        return _extract_docx(content)
    return _extract_txt(content)
