"""Export ranked candidates to CSV, Excel, or a PDF report."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
)

COLUMNS = [
    "Rank", "Name", "Email", "Phone", "Years Exp", "Match %",
    "Recommendation", "Top Skills", "Missing Skills", "Notes",
]


def _row(rank: int, c) -> list:
    s = c.score
    return [
        rank, c.name, c.email, c.phone, f"{c.years_experience:.0f}",
        f"{s.overall:.0f}" if s else "",
        s.recommendation if s else "",
        ", ".join((s.matched_skills or [])[:6]) if s else "",
        ", ".join((s.missing_skills or [])[:6]) if s else "",
        (s.recommendation_reason if s else ""),
    ]


def to_csv(ranked: list) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    for rank, c in enumerate(ranked, start=1):
        w.writerow(_row(rank, c))
    return buf.getvalue().encode("utf-8")


def to_excel(ranked: list, job_title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    ws.append([f"Resume Ranking — {job_title}"])
    ws.append([])
    ws.append(COLUMNS)
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for rank, c in enumerate(ranked, start=1):
        ws.append(_row(rank, c))
    for col in ws.columns:
        width = max((len(str(cell.value)) for cell in col if cell.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 50)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(ranked: list, job_title: str) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=letter, title="Resume Ranking Report")
    styles = getSampleStyleSheet()
    flow = [
        Paragraph("Omni Control Technology — Candidate Ranking", styles["Title"]),
        Paragraph(f"Role: {job_title}", styles["Heading2"]),
        Spacer(1, 12),
    ]
    data = [["#", "Name", "Match", "Rec.", "Email"]]
    for rank, c in enumerate(ranked, start=1):
        s = c.score
        data.append([
            rank, c.name, f"{s.overall:.0f}%" if s else "-",
            s.recommendation if s else "-", c.email,
        ])
    table = Table(data, repeatRows=1, colWidths=[24, 130, 50, 110, 160])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    flow.append(table)
    flow.append(Spacer(1, 18))

    for rank, c in enumerate(ranked, start=1):
        s = c.score
        if not s:
            continue
        flow.append(Paragraph(f"#{rank} — {c.name} ({s.overall:.0f}%)", styles["Heading3"]))
        flow.append(Paragraph(s.summary or "", styles["BodyText"]))
        if s.strengths:
            flow.append(Paragraph("<b>Strengths:</b> " + "; ".join(s.strengths), styles["BodyText"]))
        if s.weaknesses:
            flow.append(Paragraph("<b>Weaknesses:</b> " + "; ".join(s.weaknesses), styles["BodyText"]))
        flow.append(Paragraph("<b>Recommendation:</b> " + s.recommendation_reason, styles["BodyText"]))
        flow.append(Spacer(1, 10))

    doc.build(flow)
    return out.getvalue()
